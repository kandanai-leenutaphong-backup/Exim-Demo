#!/usr/bin/env python3
"""Generate the deterministic fictional ClaimPilot demonstration data pack."""

from __future__ import annotations

import csv
import io
import json
import random
from datetime import date, datetime, timezone
from pathlib import Path

import fitz
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageEnhance, ImageFilter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
PACK_DIR = DATA_DIR / "claim_pack"
CASE_ID = "ECI-2026-000125"
WATERMARK = "SAMPLE \u2014 NOT REAL DATA"
FIXED_SCREENING_TIME = "2026-08-22T03:30:17Z"

SELLER = "Example Crest Industries Ltd."
BUYER = "Zantara Meridian Trading Co."
DESTINATION = "Port of Zantara"

LINE_ITEMS = [
    ("CP-AX14", "Fictional industrial sensor modules", 300, 450.00, 135000.00),
    ("CP-BR22", "Fictional control relay assemblies", 200, 300.00, 60000.00),
    ("CP-CN08", "Fictional shielded connector kits", 500, 60.00, 30000.00),
]


FONT_FILE = Path("/System/Library/Fonts/Supplemental/Arial.ttf")
if FONT_FILE.exists():
    pdfmetrics.registerFont(TTFont("ClaimPilotArial", str(FONT_FILE)))
    FONT = "ClaimPilotArial"
else:
    FONT = "Helvetica"


def watermark_page(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFillColor(colors.Color(0.68, 0.68, 0.68, alpha=0.34))
    canvas.setFont(FONT, 31)
    canvas.translate(A4[0] / 2, A4[1] / 2)
    canvas.rotate(35)
    canvas.drawCentredString(0, 0, WATERMARK)
    canvas.restoreState()


def styles():
    base = getSampleStyleSheet()
    base["Normal"].fontName = FONT
    base["Normal"].fontSize = 9.5
    base["Normal"].leading = 13
    base["Title"].fontName = FONT
    base["Title"].fontSize = 18
    base["Heading2"].fontName = FONT
    base["Heading2"].fontSize = 12
    base.add(
        ParagraphStyle(
            "Small",
            parent=base["Normal"],
            fontName=FONT,
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#334155"),
        )
    )
    return base


def table(rows, widths=None, header=True):
    result = Table(rows, colWidths=widths, repeatRows=1 if header else 0)
    commands = [
        ("FONTNAME", (0, 0), (-1, -1), FONT),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("LEADING", (0, 0), (-1, -1), 11),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#94A3B8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]
    if header:
        commands.extend(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173B57")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), FONT),
            ]
        )
    result.setStyle(TableStyle(commands))
    return result


def build_pdf(
    path: Path,
    title: str,
    story: list,
    metadata_title: str | None = None,
    include_watermark: bool = True,
) -> None:
    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=metadata_title or title,
        author="ClaimPilot Synthetic Data Generator",
        subject="Fictional demonstration document",
    )
    s = styles()
    content = [
        Paragraph("FICTIONAL DEMONSTRATION DOCUMENT", s["Small"]),
        Spacer(1, 4 * mm),
        Paragraph(title, s["Title"]),
        Spacer(1, 5 * mm),
        *story,
    ]
    if include_watermark:
        doc.build(content, onFirstPage=watermark_page, onLaterPages=watermark_page)
    else:
        doc.build(content)


def claim_form() -> None:
    s = styles()
    rows = [
        ["Field", "Value"],
        ["Claim / Policy number", CASE_ID],
        ["Claim date", "2026-08-08"],
        ["Insured exporter", SELLER],
        ["Buyer", BUYER],
        ["Buyer country", "Republic of Zantara (fictional)"],
        ["Invoice amount", "USD 225,000.00"],
        ["Loss reason", "Payment overdue following cited shipping delay"],
        ["Status", "Submitted for demonstration review"],
    ]
    story = [
        Paragraph(
            "This form and every named party, place, transaction, and identifier are synthetic.",
            s["Normal"],
        ),
        Spacer(1, 4 * mm),
        table(rows, [55 * mm, 110 * mm]),
        Spacer(1, 8 * mm),
        Paragraph(
            "Declaration: the submitted claim package is illustrative and is not evidence of a real trade, debt, or insurance contract.",
            s["Normal"],
        ),
    ]
    build_pdf(PACK_DIR / "01_claim_form.pdf", "Trade Credit Insurance Claim Form", story)


def commercial_invoice() -> None:
    s = styles()
    rows = [["SKU", "Description", "Qty", "Unit price", "Line total"]]
    for sku, desc, qty, unit, total in LINE_ITEMS:
        rows.append([sku, desc, str(qty), f"USD {unit:,.2f}", f"USD {total:,.2f}"])
    rows.append(["", "", "", "TOTAL", "USD 225,000.00"])
    story = [
        table(
            [
                ["Invoice number", "ECI-INV-260210-125"],
                ["Invoice date", "2026-02-10"],
                ["Seller", SELLER],
                ["Buyer", BUYER],
                ["Destination", DESTINATION],
            ],
            [45 * mm, 120 * mm],
        ),
        Spacer(1, 6 * mm),
        table(rows, [22 * mm, 65 * mm, 14 * mm, 30 * mm, 34 * mm]),
        Spacer(1, 5 * mm),
        Paragraph("Payment terms: net 60 days from shipment. Currency: USD.", s["Normal"]),
    ]
    build_pdf(PACK_DIR / "02_commercial_invoice.pdf", "Commercial Invoice", story)


def packing_list() -> None:
    rows = [["SKU", "Description", "Qty", "Packages", "Net kg", "Gross kg"]]
    package_data = [(30, 540, 600), (20, 420, 470), (25, 250, 290)]
    for item, package in zip(LINE_ITEMS, package_data):
        sku, desc, qty, _, _ = item
        packages, net, gross = package
        rows.append([sku, desc, str(qty), str(packages), f"{net:,}", f"{gross:,}"])
    rows.append(["", "TOTAL", "1,000", "75", "1,210", "1,360"])
    story = [
        table(
            [
                ["Packing list", "ECI-PL-260210-125"],
                ["Related invoice", "ECI-INV-260210-125"],
                ["Exporter", SELLER],
                ["Consignee", BUYER],
            ],
            [45 * mm, 120 * mm],
        ),
        Spacer(1, 6 * mm),
        table(rows, [20 * mm, 59 * mm, 15 * mm, 20 * mm, 23 * mm, 25 * mm]),
    ]
    build_pdf(PACK_DIR / "04_packing_list.pdf", "Packing List", story)


def export_declaration_clean(path: Path) -> None:
    s = styles()
    story = [
        table(
            [
                ["Declaration reference", "ZAN-EXP-2026-0212-0041"],
                ["Exporter", SELLER],
                ["Consignee", BUYER],
                ["Export date", "2026-02-12"],
                ["Destination", DESTINATION],
                ["Declared value", "USD 225,000.00"],
                ["Packages", "75"],
                ["Country", "Republic of Zantara (fictional)"],
            ],
            [50 * mm, 115 * mm],
        ),
        Spacer(1, 8 * mm),
        Paragraph(
            "Customs note: synthetic declaration created solely for a ClaimPilot demonstration.",
            s["Normal"],
        ),
    ]
    build_pdf(path, "Export Declaration", story, include_watermark=False)


def bill_of_lading_clean(path: Path) -> None:
    s = styles()
    story = [
        table(
            [
                ["Bill of lading", "FIC-ZAN-260212-0188"],
                ["Shipper", SELLER],
                ["Consignee", BUYER],
                ["Port of loading", "Port of Northbridge (fictional)"],
                ["Port of discharge", DESTINATION],
                ["Shipment date", "2026-02-12"],
                ["Declared cargo value", "USD 225,000.00"],
                ["Packages / weight", "75 packages / 1,360 kg gross"],
                ["Carrier", "Fictional Oceanic Lines"],
            ],
            [55 * mm, 110 * mm],
        ),
        Spacer(1, 8 * mm),
        Paragraph(
            "Cargo description: industrial sensor modules, control relay assemblies, and shielded connector kits.",
            s["Normal"],
        ),
    ]
    build_pdf(path, "Bill of Lading", story, include_watermark=False)


def degrade_pdf(source: Path, destination: Path, severity: str) -> None:
    src = fitz.open(source)
    out = fitz.open()
    for page_number, page in enumerate(src):
        zoom = 1.35 if severity == "heavy" else 1.75
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        image = Image.open(io.BytesIO(pix.tobytes("png"))).convert("L")
        if severity == "heavy":
            image = image.resize(
                (int(image.width * 0.72), int(image.height * 0.72)),
                Image.Resampling.BILINEAR,
            )
            image = ImageEnhance.Contrast(image).enhance(0.76)
            image = image.filter(ImageFilter.GaussianBlur(radius=0.65))
            rng = random.Random(125 + page_number)
            pixels = image.load()
            for _ in range(2200):
                x = rng.randrange(image.width)
                y = rng.randrange(image.height)
                pixels[x, y] = rng.choice((150, 175, 205, 235))
        else:
            image = ImageEnhance.Contrast(image).enhance(0.9)
            image = image.filter(ImageFilter.GaussianBlur(radius=0.25))
        buffer = io.BytesIO()
        image.save(buffer, format="JPEG", quality=58 if severity == "heavy" else 76)
        output_page = out.new_page(width=A4[0], height=A4[1])
        output_page.insert_image(output_page.rect, stream=buffer.getvalue())
        overlay_font = "helv"
        if FONT_FILE.exists():
            output_page.insert_font(fontname="ClaimPilotArial", fontfile=str(FONT_FILE))
            overlay_font = "ClaimPilotArial"
        ocr_lines = (
            ["Shipment date 2026-02-12", "Declared cargo value USD 225,000.00"]
            if severity == "heavy"
            else ["Destination Port of Zantara", "Declared value USD 225,000.00"]
        )
        for line_number, line in enumerate(ocr_lines):
            output_page.insert_text(
                fitz.Point(12, 12 + line_number * 3),
                line,
                fontsize=1,
                fontname="helv",
                render_mode=3,
            )
        shape = output_page.new_shape()
        shape.insert_text(
            fitz.Point(110, 470),
            WATERMARK,
            fontsize=28,
            fontname=overlay_font,
            color=(0.62, 0.62, 0.62),
            rotate=0,
            fill_opacity=0.52,
        )
        shape.commit()
    out.set_metadata(
        {
            "title": destination.stem,
            "author": "ClaimPilot Synthetic Data Generator",
            "subject": "Fictional degraded demonstration scan",
        }
    )
    out.save(destination)
    out.close()
    src.close()


def policy() -> None:
    s = styles()
    page_one = [
        Paragraph("POLICY SCHEDULE", s["Heading2"]),
        table(
            [
                ["Policy number", CASE_ID],
                ["Policyholder", SELLER],
                ["Policy period", "2026-01-01 through 2026-12-31"],
                ["Maximum liability", "USD 500,000.00"],
                ["Governing territory", "Fictional demonstration jurisdictions only"],
            ],
            [50 * mm, 115 * mm],
        ),
        Spacer(1, 6 * mm),
        Paragraph(
            "This policy is synthetic. It confers no coverage, rights, or obligations.",
            s["Normal"],
        ),
        PageBreak(),
        Paragraph("CLAIMS CONDITIONS", s["Heading2"]),
        Spacer(1, 4 * mm),
        Paragraph(
            "<b>Clause 7.3 - Filing Window.</b> A claim must be filed within 180 calendar days after the shipment date shown in acceptable shipping evidence. If the shipment date cannot be verified with sufficient confidence, no filing deadline shall be calculated until reviewed.",
            s["Normal"],
        ),
        Spacer(1, 5 * mm),
        Paragraph(
            "<b>Clause 7.4 - Supporting Records.</b> The insured should provide a claim form, commercial invoice, bill of lading, packing list, export declaration, buyer correspondence, payment history, policy, proof of delivery, and power of attorney when applicable.",
            s["Normal"],
        ),
        Spacer(1, 5 * mm),
        Paragraph(
            "<b>Stable citation:</b> Policy ECI-2026-000125, Clause 7.3, page 2.",
            s["Normal"],
        ),
    ]
    build_pdf(PACK_DIR / "08_policy.pdf", "Trade Credit Insurance Policy", page_one)


def correspondence() -> None:
    text = """FICTIONAL DEMONSTRATION CORRESPONDENCE
Case: ECI-2026-000125
From: Accounts Payable, Zantara Meridian Trading Co.
To: Credit Control, Example Crest Industries Ltd.
Date: 2026-07-28

เรียนทีมงาน Example Crest,

เราได้รับเอกสารสำหรับใบแจ้งหนี้ ECI-INV-260210-125 แล้ว แต่การจัดส่งล่าช้ากว่ากำหนด
และทำให้กระบวนการรับสินค้าและอนุมัติการชำระเงินของเราล่าช้าไปด้วย

We cite the shipping delay as the reason payment has not yet been released.
This message is not a formal dispute of the invoice, goods, quantity, or amount.
We are reviewing the payment schedule and will provide an update.

ขอแสดงความนับถือ
Fictional Accounts Payable Team
"""
    (PACK_DIR / "06_correspondence.txt").write_text(text, encoding="utf-8")


def payment_history() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Payment History"
    headers = [
        "Month",
        "Invoices Due",
        "Amount Due (USD)",
        "Amount Paid (USD)",
        "Days Late",
        "Status",
    ]
    sheet.append(headers)
    rng = random.Random(2026000125)
    for index in range(36):
        year = 2023 + (index + 1) // 12
        month = (index + 1) % 12 + 1
        month_value = date(year, month, 1)
        invoices = 2 + index % 4
        amount_due = 28000 + rng.randrange(0, 18001)
        if index < 30:
            days_late = rng.randrange(0, 13)
            ratio = 1.0 if days_late < 8 else 0.96
            status = "Paid" if ratio == 1.0 else "Paid late"
        else:
            days_late = [18, 26, 39, 55, 78, 104][index - 30]
            ratio = [0.94, 0.88, 0.77, 0.61, 0.38, 0.10][index - 30]
            status = "Deteriorating" if index < 35 else "Severely overdue"
        amount_paid = round(amount_due * ratio, 2)
        sheet.append([month_value, invoices, amount_due, amount_paid, days_late, status])

    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173B57")
        cell.alignment = Alignment(horizontal="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", color="0000FF")
    for row_number in range(32, 38):
        for cell in sheet[row_number]:
            cell.fill = PatternFill("solid", fgColor="FDE68A")
    for cell in sheet["A"][1:]:
        cell.number_format = "yyyy-mm"
    for column in ("C", "D"):
        for cell in sheet[column][1:]:
            cell.number_format = '$#,##0.00;($#,##0.00);-'
    widths = {"A": 14, "B": 14, "C": 21, "D": 22, "E": 12, "F": 20}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:F37"
    workbook.properties.title = "Fictional ClaimPilot Payment History"
    workbook.properties.description = "Synthetic data for case ECI-2026-000125"
    workbook.save(PACK_DIR / "07_payment_history.xlsx")


def history() -> None:
    fields = [
        "claim_id",
        "claim_amount_usd",
        "document_completeness",
        "max_days_late",
        "buyer_risk_score",
        "shipment_confidence",
        "sanctions_match",
        "was_accepted",
    ]
    rng = random.Random(125)
    with (DATA_DIR / "history.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for index in range(1, 201):
            completeness = rng.randint(5, 10)
            max_days_late = rng.randint(0, 150)
            buyer_risk = rng.randint(18, 92)
            shipment_confidence = round(rng.uniform(0.45, 0.99), 2)
            sanctions_match = 1 if index % 37 == 0 else 0
            accepted = int(
                completeness >= 8
                and max_days_late <= 105
                and buyer_risk <= 72
                and shipment_confidence >= 0.60
                and not sanctions_match
            )
            writer.writerow(
                {
                    "claim_id": f"SYN-2026-{index:06d}",
                    "claim_amount_usd": rng.randrange(15000, 450001, 500),
                    "document_completeness": completeness,
                    "max_days_late": max_days_late,
                    "buyer_risk_score": buyer_risk,
                    "shipment_confidence": f"{shipment_confidence:.2f}",
                    "sanctions_match": sanctions_match,
                    "was_accepted": accepted,
                }
            )


def sanctions() -> None:
    payload = {
        "label": "ILLUSTRATIVE \u2014 NOT AN OFFICIAL SANCTIONS LIST",
        "version": "2026.08-demo",
        "generated_for": "ClaimPilot synthetic demonstration",
        "entries": [
            {
                "entry_id": "ILL-ZAN-001",
                "entity_name": BUYER,
                "aliases": ["Zantara Meridian Trading"],
                "country": "Republic of Zantara",
                "program": "Fictional Zantara Trade Review",
            },
            {
                "entry_id": "ILL-NOV-002",
                "entity_name": "Novara Harbor Components Cooperative",
                "aliases": [],
                "country": "Federation of Novara",
                "program": "Fictional Novara Demonstration Program",
            },
        ],
    }
    (DATA_DIR / "sanctions.json").write_text(
        json.dumps(payload, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )


def cached_run() -> None:
    thai_report = (
        "คำขอเคลมกรณี ECI-2026-000125 มีเอกสารครบ 8 จาก 10 รายการ "
        "ยอดตามใบแจ้งหนี้คือ USD 225,000.00 และผู้ซื้อระบุว่าการชำระเงินล่าช้าเพราะการขนส่งล่าช้า โดยไม่มีข้อพิพาทอย่างเป็นทางการ\n\n"
        "วันที่จัดส่งที่อ่านจาก 03_bill_of_lading.pdf คือ 2026-02-12 แต่ค่าความเชื่อมั่นต่ำกว่าเกณฑ์ 0.60 "
        "ดังนั้นกำหนดยื่นภายใน 180 วันตาม Policy ECI-2026-000125, Clause 7.3, page 2 ยังไม่ได้รับการยืนยันและไม่มีการคำนวณวันครบกำหนด\n\n"
        "ผลคัดกรองแบบกำหนดตายตัวพบรายการที่ตรงกับ Zantara Meridian Trading Co. ในบัญชีตัวอย่าง และคะแนนเคสคือ 42 จาก 100 "
        "ข้อมูลทั้งหมดเป็นข้อมูลสังเคราะห์สำหรับการสาธิต"
    )
    panels = [
        {
            "panel": "classification",
            "status": "complete",
            "timestamp": FIXED_SCREENING_TIME,
            "data": {
                "documents": [
                    {"file": "01_claim_form.pdf", "type": "claim_form", "confidence": 0.99},
                    {"file": "02_commercial_invoice.pdf", "type": "commercial_invoice", "confidence": 0.99},
                    {"file": "03_bill_of_lading.pdf", "type": "bill_of_lading", "confidence": 0.91},
                    {"file": "04_packing_list.pdf", "type": "packing_list", "confidence": 0.98},
                    {"file": "05_export_declaration.pdf", "type": "export_declaration", "confidence": 0.93},
                    {"file": "06_correspondence.txt", "type": "correspondence", "confidence": 0.99},
                    {"file": "07_payment_history.xlsx", "type": "payment_history", "confidence": 0.99},
                    {"file": "08_policy.pdf", "type": "policy", "confidence": 0.99},
                ],
                "completeness": {
                    "present": 8,
                    "required": 10,
                    "ratio": 0.8,
                    "missing": ["power_of_attorney", "proof_of_delivery"],
                },
            },
        },
        {
            "panel": "extraction",
            "status": "warning",
            "timestamp": FIXED_SCREENING_TIME,
            "data": {
                "documents": [
                    {
                        "file": "01_claim_form.pdf",
                        "fields": [
                            {"name": "claim_date", "value": "2026-08-08", "confidence": 0.99},
                            {"name": "policy_number", "value": CASE_ID, "confidence": 0.99},
                            {"name": "buyer", "value": BUYER, "confidence": 0.98},
                        ],
                    },
                    {
                        "file": "02_commercial_invoice.pdf",
                        "fields": [
                            {"name": "invoice_date", "value": "2026-02-10", "confidence": 0.98},
                            {"name": "invoice_amount", "value": "USD 225,000.00", "confidence": 0.99},
                        ],
                    },
                    {
                        "file": "03_bill_of_lading.pdf",
                        "fields": [
                            {"name": "shipment_date", "value": "2026-02-12", "confidence": 0.54},
                            {"name": "declared_cargo_value", "value": "USD 225,000.00", "confidence": 0.93},
                        ],
                    },
                    {
                        "file": "05_export_declaration.pdf",
                        "fields": [
                            {"name": "destination_port", "value": DESTINATION, "confidence": 0.91},
                        ],
                    },
                ]
            },
        },
        {
            "panel": "language",
            "status": "complete",
            "timestamp": FIXED_SCREENING_TIME,
            "data": {
                "detected_languages": [
                    {"language": "Thai", "iso6391": "th", "confidence": 0.99},
                    {"language": "English", "iso6391": "en", "confidence": 0.99},
                ],
                "entities": [
                    {"text": BUYER, "category": "Organization"},
                    {"text": "shipping delay", "category": "Event"},
                ],
                "formal_dispute": False,
                "summary_en": "Buyer cites a shipping delay; no formal dispute is present.",
                "summary_th": "ผู้ซื้ออ้างถึงความล่าช้าในการขนส่ง และไม่พบข้อพิพาทอย่างเป็นทางการ",
            },
        },
        {
            "panel": "policy",
            "status": "warning",
            "timestamp": FIXED_SCREENING_TIME,
            "data": {
                "answer": "The policy requires filing within 180 calendar days of shipment, but this case cannot be evaluated because the shipment date is unverified.",
                "citation": {
                    "label": "Clause 7.3, page 2",
                    "clause": "7.3",
                    "page": 2,
                    "file": "08_policy.pdf",
                },
                "filing_window": {
                    "days": 180,
                    "status": "unverified",
                    "deadline": None,
                    "source_field": "shipment_date",
                    "source_file": "03_bill_of_lading.pdf",
                    "reason": "shipment_date confidence 0.54 is below the 0.60 confidence gate",
                },
            },
        },
        {
            "panel": "screening",
            "status": "warning",
            "timestamp": FIXED_SCREENING_TIME,
            "data": {
                "result": "refer",
                "counterparty": BUYER,
                "matched_entry": {
                    "entry_id": "ILL-ZAN-001",
                    "entity_name": BUYER,
                    "country": "Republic of Zantara",
                },
                "list_label": "ILLUSTRATIVE \u2014 NOT AN OFFICIAL SANCTIONS LIST",
                "list_version": "2026.08-demo",
                "checked_at": FIXED_SCREENING_TIME,
                "method": "deterministic exact-name lookup",
                "uses_llm": False,
            },
        },
        {
            "panel": "score",
            "status": "complete",
            "timestamp": FIXED_SCREENING_TIME,
            "data": {
                "score": 42,
                "scale": 100,
                "disposition": "manual_review",
                "top_factors": [
                    {"factor": "Document completeness", "value": "8/10", "impact": "negative"},
                    {"factor": "Payment aging trend", "value": "worsening final six months", "impact": "negative"},
                    {"factor": "Prior claim count", "value": 1, "impact": "neutral"},
                ],
                "unverified_inputs": ["shipment_date from 03_bill_of_lading.pdf"],
            },
        },
        {
            "panel": "report",
            "status": "complete",
            "timestamp": FIXED_SCREENING_TIME,
            "data": {
                "language": "th",
                "paragraph_count": 3,
                "thai_report": thai_report,
                "input_panels": ["classification", "extraction", "language", "policy", "screening", "score"],
                "grounding_rule": "No facts outside panel results 1-6.",
            },
        },
    ]
    payload = {
        "schema_version": "1.0",
        "case_id": CASE_ID,
        "generated_at": FIXED_SCREENING_TIME,
        "panels": panels,
    }
    (DATA_DIR / "cached_run.json").write_text(
        json.dumps(payload, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )


def main() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    PACK_DIR.mkdir(parents=True, exist_ok=True)

    claim_form()
    commercial_invoice()
    packing_list()
    correspondence()
    payment_history()
    policy()

    temp_bill = PACK_DIR / ".bill_clean.pdf"
    temp_export = PACK_DIR / ".export_clean.pdf"
    bill_of_lading_clean(temp_bill)
    export_declaration_clean(temp_export)
    degrade_pdf(temp_bill, PACK_DIR / "03_bill_of_lading.pdf", "heavy")
    degrade_pdf(temp_export, PACK_DIR / "05_export_declaration.pdf", "light")
    temp_bill.unlink()
    temp_export.unlink()

    history()
    sanctions()
    cached_run()
    print(f"Generated deterministic ClaimPilot data under {DATA_DIR}")


if __name__ == "__main__":
    main()
