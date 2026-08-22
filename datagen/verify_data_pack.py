#!/usr/bin/env python3
"""Verify the generated ClaimPilot pack and deterministic fixtures."""

from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

import fitz
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
PACK_DIR = DATA_DIR / "claim_pack"
WATERMARK = "SAMPLE \u2014 NOT REAL DATA"

EXPECTED_PACK_FILES = {
    "01_claim_form.pdf",
    "02_commercial_invoice.pdf",
    "03_bill_of_lading.pdf",
    "04_packing_list.pdf",
    "05_export_declaration.pdf",
    "06_correspondence.txt",
    "07_payment_history.xlsx",
    "08_policy.pdf",
}
EXPECTED_PDFS = {name for name in EXPECTED_PACK_FILES if name.endswith(".pdf")}


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def extracted_text(path: Path) -> str:
    with fitz.open(path) as document:
        return "\n".join(page.get_text() for page in document)


def verify_watermarks() -> None:
    for name in sorted(EXPECTED_PDFS):
        path = PACK_DIR / name
        with fitz.open(path) as document:
            require(document.page_count > 0, f"{name} has no pages")
            for page_number, page in enumerate(document, start=1):
                text = page.get_text().replace("\u00a0", " ")
                require(
                    WATERMARK in text,
                    f"{name} page {page_number} lacks exact watermark",
                )


def verify_documents() -> None:
    actual = {path.name for path in PACK_DIR.iterdir() if not path.name.startswith(".")}
    require(actual == EXPECTED_PACK_FILES, f"pack manifest mismatch: {sorted(actual)}")
    require(len(EXPECTED_PDFS) == 6, "enumerated pack must contain six PDFs")
    require(len(EXPECTED_PACK_FILES) == 8, "claim pack must contain eight submitted items")
    verify_watermarks()

    claim = extracted_text(PACK_DIR / "01_claim_form.pdf")
    require("2026-08-08" in claim and "ECI-2026-000125" in claim, "claim values")
    invoice = extracted_text(PACK_DIR / "02_commercial_invoice.pdf")
    require("2026-02-10" in invoice and "USD 225,000.00" in invoice, "invoice values")
    bill = extracted_text(PACK_DIR / "03_bill_of_lading.pdf")
    require("2026-02-12" in bill and "USD 225,000.00" in bill, "bill values")
    packing = extracted_text(PACK_DIR / "04_packing_list.pdf")
    for sku, qty in (("CP-AX14", "300"), ("CP-BR22", "200"), ("CP-CN08", "500")):
        require(sku in packing and qty in packing, f"packing line item {sku}")
    export = extracted_text(PACK_DIR / "05_export_declaration.pdf")
    require("Port of Zantara" in export, "export destination")
    policy = extracted_text(PACK_DIR / "08_policy.pdf")
    require("180 calendar days" in policy, "policy filing window")
    require(
        "Policy ECI-2026-000125, Clause 7.3, page 2" in policy,
        "stable policy citation",
    )
    correspondence = (PACK_DIR / "06_correspondence.txt").read_text(encoding="utf-8")
    require("shipping delay" in correspondence, "correspondence shipping delay")
    require("not a formal dispute" in correspondence, "correspondence dispute status")

    workbook = load_workbook(PACK_DIR / "07_payment_history.xlsx", data_only=False)
    sheet = workbook["Payment History"]
    require(sheet.max_row == 37, "payment history must have 36 monthly rows")
    days_late = [sheet.cell(row=row, column=5).value for row in range(32, 38)]
    require(days_late == sorted(days_late), "final six months must worsen")
    require(days_late == [18, 26, 39, 55, 78, 104], "final six fixture mismatch")


def verify_data_files() -> None:
    with (DATA_DIR / "history.csv").open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    require(len(rows) == 200, "history.csv must contain 200 rows")
    require("was_accepted" in rows[0], "history.csv lacks was_accepted")

    sanctions = json.loads((DATA_DIR / "sanctions.json").read_text(encoding="utf-8"))
    require(
        sanctions["label"] == "ILLUSTRATIVE \u2014 NOT AN OFFICIAL SANCTIONS LIST",
        "sanctions label mismatch",
    )
    require(
        all(
            "fictional" in entry["country"].lower()
            or entry["country"] in {"Republic of Zantara", "Federation of Novara"}
            for entry in sanctions["entries"]
        ),
        "sanctions countries must be fictional",
    )

    cached = json.loads((DATA_DIR / "cached_run.json").read_text(encoding="utf-8"))
    panels = cached["panels"]
    require(len(panels) == 7, "cached run must have seven panels")
    require(
        [panel["panel"] for panel in panels]
        == ["classification", "extraction", "language", "policy", "screening", "score", "report"],
        "panel order",
    )
    completeness = panels[0]["data"]["completeness"]
    require(
        completeness["present"] == 8 and completeness["required"] == 10,
        "completeness must be 8/10",
    )
    require(
        completeness["missing"] == ["power_of_attorney", "proof_of_delivery"],
        "missing document fixture mismatch",
    )
    extraction = panels[1]
    shipment = next(
        field
        for document in extraction["data"]["documents"]
        for field in document["fields"]
        if field["name"] == "shipment_date"
    )
    require(
        shipment["confidence"] < 0.60,
        "shipment confidence must remain below 0.60",
    )
    policy = panels[3]
    require(
        policy["data"]["filing_window"]["status"] == "unverified"
        and policy["data"]["filing_window"]["deadline"] is None,
        "policy confidence gate mismatch",
    )
    screening = panels[4]
    require(
        screening["data"]["uses_llm"] is False
        and screening["data"]["result"] == "refer"
        and screening["data"]["checked_at"] == "2026-08-22T03:30:17Z",
        "screening fixture mismatch",
    )
    report = panels[6]["data"]["thai_report"]
    require(len(report.split("\n\n")) == 3, "Thai report must have three paragraphs")
    require("03_bill_of_lading.pdf" in report, "Thai report must name source file")
    require("ยังไม่ได้รับการยืนยัน" in report, "Thai report must state unverified deadline")


def main() -> int:
    try:
        verify_documents()
        verify_data_files()
    except (AssertionError, FileNotFoundError, KeyError) as error:
        print(f"FAIL: {error}", file=sys.stderr)
        return 1
    print(
        "PASS: 8-item claim pack (6 PDFs + TXT + XLSX), all PDF pages watermarked; "
        "200 history rows; cached confidence gate enforced."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
